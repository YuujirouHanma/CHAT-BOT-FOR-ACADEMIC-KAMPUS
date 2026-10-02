"""Document parser using Unstructured.io.

Supports: PDF, DOCX, PPTX, TXT, MD, CSV, XLSX, HTML, RTF, RST, EPUB, TSV.
Unstructured's partition() auto-detects file type.

Strategy:
- PDF: 'fast' (pdfminer, no OCR/Poppler needed — see feedback-parser-strategy).
- Other formats (docx, pptx, xlsx, ...): let partition() auto-route to the
  format-specific partitioner with its own defaults. We deliberately do NOT
  pass infer_table_structure: partition() already forwards it internally, so
  passing it again raises "got multiple values for keyword argument".
"""
from __future__ import annotations

import base64
import uuid
from pathlib import Path
from typing import Any

try:
    from pdf2image import convert_from_path
except Exception:  # pragma: no cover - optional runtime dependency
    convert_from_path = None

try:
    import pytesseract
except Exception:  # pragma: no cover - optional runtime dependency
    pytesseract = None

from unstructured.documents.elements import (
    Element,
)
from unstructured.documents.elements import (
    Image as UImage,
)
from unstructured.documents.elements import (
    Table as UTable,
)
from unstructured.partition.auto import partition

from src.ingestion.validators import validate_indexable
from src.schemas import ElementType, ParsedElement
from src.utils.logger import logger

_PDF_EXTENSIONS = {"pdf"}

# Below this many extracted text chars, a PDF is treated as "no text" — almost
# always a scanned/image-based PDF whose text lives in images. `fast` (pdfminer)
# does no OCR, so we retry with OCR. See feedback-parser-strategy.
_MIN_PDF_TEXT_CHARS = 20

# Ambang per halaman, bukan hanya total. Kasus yang lolos dari ambang total:
# slide kuliah yang lapisan teksnya cuma berisi judul sementara isinya gambar.
# Terukur pada rka_dasprog_tm01.pdf — 29 halaman, 719 karakter (24,8/halaman)
# lewat pdfminer versus 21.441 karakter lewat OCR. Dengan ambang total 20
# karakter berkas itu dianggap sehat, dan 97% isinya tidak pernah masuk indeks.
#
# 100 dipilih jauh di bawah PDF yang benar-benar berisi teks (terukur 256–1787
# karakter/halaman) supaya OCR yang mahal (~4 detik/halaman) hanya berjalan saat
# lapisan teksnya memang nyaris kosong. Hasil OCR tetap hanya dipakai bila lebih
# banyak dari hasil fast, jadi ambang ini tidak pernah bisa memperburuk.
_MIN_PDF_CHARS_PER_PAGE = 100

# Padanan untuk slide. Deck yang seluruhnya gambar melewati partisi tanpa galat
# sekaligus tanpa hasil — terukur pada rka_dasprog_PR01: 8 slide, 10 gambar,
# NOL karakter, 0 elemen. Berkasnya tercatat berhasil diunggah lalu tidak pernah
# bisa ditemukan AI. Ambangnya rendah karena hanya untuk membedakan "benar-benar
# tidak ada teks" dari "ada teks meski sedikit"; deck yang punya teks tidak
# disentuh supaya tidak membayar satu panggilan LLM per gambar untuk isi yang
# sudah terwakili teksnya.
_MIN_PPTX_TEXT_CHARS = 20

# Ambang per slide, bukan hanya total — pelajaran yang sama dari sisi PDF.
# Ambang total saja meloloskan bentuk slide kuliah yang PALING umum: tiap slide
# punya judul singkat sementara seluruh isinya gambar. Deck seperti itu
# menghasilkan ratusan karakter secara total sehingga dianggap sehat, padahal
# yang terindeks hanya deretan judul.
#
# 120 dipilih dari pengukuran: deck yang benar-benar berisi teks menghasilkan
# 486 karakter/slide, sementara deck yang hanya berjudul menghasilkan sekitar
# 31. Jurang di antara keduanya lebar, jadi ambang ini tidak sensitif terhadap
# pilihan angka persisnya.
_MIN_PPTX_CHARS_PER_SLIDE = 120

# Logo, ikon, dan hiasan bullet ikut terhitung sebagai PICTURE. Masing-masing
# akan memicu satu panggilan LLM dan menghasilkan deskripsi yang tidak seorang
# pun cari, jadi yang di bawah ambang ini dibuang.
_MIN_SLIDE_IMAGE_BYTES = 8192

# Batas atas gambar per berkas. Tiap gambar berbiaya satu panggilan LLM saat
# indexing, jadi tanpa batas ini satu deck ratusan slide bisa membuat sekali
# unggah menghabiskan waktu dan biaya yang tidak diperkirakan dosen.
_MAX_SLIDE_IMAGES = 40


def parse_document(
    file_path: Path,
    content_id: str | None = None,
) -> list[ParsedElement]:
    """Parse a document into a list of ParsedElement.

    Args:
        file_path: Path to the file on disk.
        content_id: Content identifier to stamp on every element.
    """
    validate_indexable(file_path)
    logger.info("Parsing document: {}", file_path.name)

    parsed = _partition_and_assemble(
        file_path, _build_partition_kwargs(file_path), content_id
    )

    # OCR fallback: a scanned/image PDF parses without error but yields ~no text.
    # Retry with strategy="ocr_only" so it becomes searchable instead of empty.
    if _is_pdf(file_path) and _pdf_text_layer_too_thin(file_path, parsed):
        ocr_parsed = _try_ocr(file_path, content_id)
        if ocr_parsed is not None and _text_char_count(ocr_parsed) > _text_char_count(parsed):
            logger.info(
                "OCR dipakai untuk '{}': {} → {} karakter",
                file_path.name, _text_char_count(parsed), _text_char_count(ocr_parsed),
            )
            parsed = ocr_parsed

    # Padanan OCR untuk slide: deck yang seluruhnya gambar melewati partisi
    # tanpa galat sekaligus tanpa hasil, jadi gambarnya diambil langsung dan
    # dideskripsikan peringkas multimodal.
    if _is_pptx(file_path) and _slide_images_needed(parsed, file_path):
        gambar_slide = _extract_pptx_images(file_path, content_id)
        if gambar_slide:
            logger.info(
                "Gambar slide dipakai untuk '{}': {} gambar diambil",
                file_path.name, len(gambar_slide),
            )
            parsed = parsed + gambar_slide

    counts = _count_by_type(parsed)
    logger.info(
        "Parsed {}: {} text, {} tables, {} images",
        file_path.name,
        counts[ElementType.TEXT],
        counts[ElementType.TABLE],
        counts[ElementType.IMAGE],
    )
    return parsed


def _partition_and_assemble(
    file_path: Path, partition_kwargs: dict[str, Any], content_id: str | None,
) -> list[ParsedElement]:
    """Run partition() with the given kwargs and assemble ParsedElements.

    A partition failure is fatal (wrapped in RuntimeError) — the caller cannot
    recover from a corrupt/unreadable file.
    """
    try:
        elements: list[Element] = partition(filename=str(file_path), **partition_kwargs)
    except Exception as exc:
        logger.exception("Partition failed for {}", file_path.name)
        raise RuntimeError(f"Failed to parse {file_path.name}: {exc}") from exc

    return _assemble_elements(
        elements, source_file=file_path.name, content_id=content_id,
    )


def _try_ocr(file_path: Path, content_id: str | None) -> list[ParsedElement] | None:
    """Re-parse a PDF with OCR. Returns None (not raising) when OCR is
    unavailable — e.g. Tesseract/Poppler not installed — so a scanned PDF
    degrades to an empty parse instead of crashing the upload.
    """
    parsed: list[ParsedElement] | None = None
    try:
        elements: list[Element] = partition(
            filename=str(file_path),
            strategy="ocr_only",
            languages=["ind", "eng"],
        )
    except Exception as exc:
        logger.warning(
            "OCR fallback unavailable for {} ({}). To read scanned PDFs, install "
            "Tesseract (tesseract-ocr + tesseract-ocr-ind) and Poppler in the "
            "runtime image.",
            file_path.name,
            exc,
        )
    else:
        parsed = _assemble_elements(
            elements, source_file=file_path.name, content_id=content_id,
        )

    local_parsed = _try_local_pdf_ocr(file_path, content_id)
    if local_parsed is not None and (
        parsed is None or _text_char_count(local_parsed) > _text_char_count(parsed)
    ):
        return local_parsed

    return parsed


def _try_local_pdf_ocr(
    file_path: Path,
    content_id: str | None,
) -> list[ParsedElement] | None:
    """Fallback OCR path that uses local Tesseract binaries directly.

    This keeps scanned PDFs searchable even when Unstructured's OCR strategy is
    unavailable in the runtime image.
    """
    if convert_from_path is None or pytesseract is None:
        logger.warning(
            "Local OCR fallback unavailable for {}. Install pdf2image and "
            "pytesseract alongside Tesseract and Poppler in the runtime image.",
            file_path.name,
        )
        return None

    try:
        pages = convert_from_path(str(file_path))
    except Exception as exc:
        logger.warning(
            "Local OCR fallback failed for {} ({}).", file_path.name, exc,
        )
        return None

    parsed: list[ParsedElement] = []
    for page_number, image in enumerate(pages, start=1):
        try:
            text = pytesseract.image_to_string(image, lang="ind+eng").strip()
        except Exception as exc:
            logger.warning(
                "Local OCR failed on page {} of {} ({}).",
                page_number,
                file_path.name,
                exc,
            )
            continue

        if text:
            parsed.append(
                _make_text_element([text], file_path.name, page_number, content_id)
            )

    return parsed


def _is_pdf(file_path: Path) -> bool:
    return file_path.suffix.lower().lstrip(".") in _PDF_EXTENSIONS


def _is_pptx(file_path: Path) -> bool:
    return file_path.suffix.lower() == ".pptx"


def _slide_images_needed(parsed: list[ParsedElement], file_path: Path) -> bool:
    """Apakah deck ini perlu diambil gambarnya sendiri?

    Dua ukuran, bukan satu. Ambang total menangkap deck yang benar-benar tanpa
    teks; ambang per slide menangkap bentuk yang jauh lebih umum — tiap slide
    berjudul singkat sementara seluruh isinya gambar, yang secara total terlihat
    sehat padahal yang terindeks hanya deretan judul.

    Deck yang sudah punya teks memadai tidak disentuh: menambahkan deskripsi
    gambar di sana berarti satu panggilan LLM per gambar untuk isi yang sudah
    terwakili teksnya.
    """
    if any(e.element_type == ElementType.IMAGE for e in parsed):
        return False

    jumlah = _text_char_count(parsed)
    if jumlah < _MIN_PPTX_TEXT_CHARS:
        logger.warning(
            "Slide '{}' hanya menghasilkan {} karakter dan nol gambar — "
            "mencoba mengambil gambarnya langsung dari berkas",
            file_path.name, jumlah,
        )
        return True

    slide = _pptx_slide_count(file_path)
    if slide and jumlah / slide < _MIN_PPTX_CHARS_PER_SLIDE:
        logger.warning(
            "Slide '{}' hanya {:.0f} karakter/slide ({} karakter, {} slide) — "
            "isinya kemungkinan besar gambar; mencoba mengambilnya langsung",
            file_path.name, jumlah / slide, jumlah, slide,
        )
        return True

    return False


def _pptx_slide_count(file_path: Path) -> int | None:
    """Jumlah slide, atau None bila tidak terbaca."""
    try:
        from pptx import Presentation

        return len(Presentation(str(file_path)).slides)
    except Exception as exc:  # noqa: BLE001 - berkas rusak jangan menggagalkan parse
        logger.debug("Tidak bisa membaca jumlah slide {}: {}", file_path.name, exc)
        return None


def _extract_pptx_images(
    file_path: Path, content_id: str | None
) -> list[ParsedElement]:
    """Ambil gambar dari tiap slide sebagai elemen IMAGE.

    Dibutuhkan karena deck hasil pindai atau ekspor gambar melewati partisi
    normal tanpa galat sekaligus tanpa hasil: terukur pada rka_dasprog_PR01,
    8 slide berisi 10 gambar dan NOL karakter teks menghasilkan 0 elemen.
    Berkasnya tercatat berhasil diunggah, lalu tidak pernah bisa ditemukan AI.

    Gambar yang dihasilkan di sini akan dideskripsikan peringkas multimodal
    (lihat indexing/summarizer.py), dan deskripsi itulah yang masuk indeks.
    """
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError as exc:
        logger.warning("python-pptx tidak tersedia, gambar slide dilewati: {}", exc)
        return []

    def kumpulkan(bentuk_bentuk, nomor_slide: int) -> list[ParsedElement]:
        hasil: list[ParsedElement] = []
        for bentuk in bentuk_bentuk:
            # Gambar sering dibungkus dalam grup, dan grup bisa bersarang.
            if bentuk.shape_type == MSO_SHAPE_TYPE.GROUP:
                hasil.extend(kumpulkan(bentuk.shapes, nomor_slide))
                continue
            if bentuk.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue

            try:
                blob = bentuk.image.blob
            except Exception as exc:  # noqa: BLE001 - satu gambar rusak jangan menggagalkan berkas
                logger.debug("Gambar di slide {} tidak terbaca: {}", nomor_slide, exc)
                continue

            # Logo, ikon, dan hiasan bullet lolos sebagai PICTURE juga. Masing-
            # masing memicu satu panggilan LLM dan menghasilkan deskripsi yang
            # tidak seorang pun cari, jadi yang kecil dibuang.
            if len(blob) < _MIN_SLIDE_IMAGE_BYTES:
                continue

            hasil.append(
                ParsedElement(
                    element_id=str(uuid.uuid4()),
                    element_type=ElementType.IMAGE,
                    content="",
                    image_base64=base64.b64encode(blob).decode("ascii"),
                    source_file=file_path.name,
                    page_number=nomor_slide,
                    content_id=content_id,
                )
            )
        return hasil

    try:
        presentasi = Presentation(str(file_path))
    except Exception as exc:  # noqa: BLE001 - berkas rusak jangan menggagalkan parse
        logger.warning("Tidak bisa membuka '{}' sebagai pptx: {}", file_path.name, exc)
        return []

    gambar: list[ParsedElement] = []
    for nomor, slide in enumerate(presentasi.slides, start=1):
        gambar.extend(kumpulkan(slide.shapes, nomor))

    # Tiap gambar berbiaya satu panggilan LLM saat indexing. Deck ratusan slide
    # bisa membuat satu unggahan menghabiskan waktu dan biaya tak terduga.
    if len(gambar) > _MAX_SLIDE_IMAGES:
        logger.warning(
            "'{}' punya {} gambar, dipotong ke {} teratas",
            file_path.name, len(gambar), _MAX_SLIDE_IMAGES,
        )
        gambar = gambar[:_MAX_SLIDE_IMAGES]

    return gambar


def _text_char_count(parsed: list[ParsedElement]) -> int:
    return sum(len(e.content) for e in parsed if e.element_type == ElementType.TEXT)


def _pdf_page_count(file_path: Path) -> int | None:
    """Jumlah halaman PDF, atau None bila tidak terbaca."""
    try:
        from pypdf import PdfReader

        return len(PdfReader(str(file_path)).pages)
    except Exception as exc:  # noqa: BLE001 - PDF rusak jangan menggagalkan parse
        logger.debug("Tidak bisa membaca jumlah halaman {}: {}", file_path.name, exc)
        return None


def _pdf_text_layer_too_thin(file_path: Path, parsed: list[ParsedElement]) -> bool:
    """Apakah PDF ini perlu dicoba ulang dengan OCR?

    Dua ukuran, bukan satu. Ambang total menangkap PDF hasil pindai yang benar-
    benar kosong; ambang per halaman menangkap slide kuliah yang lapisan teksnya
    hanya berisi judul — kasus yang lolos dari ambang total karena judul saja
    sudah melewati 20 karakter, padahal isinya seluruhnya gambar.
    """
    jumlah = _text_char_count(parsed)
    if jumlah < _MIN_PDF_TEXT_CHARS:
        logger.warning(
            "PDF '{}' hanya menghasilkan {} karakter — kemungkinan hasil pindai; "
            "mencoba ulang dengan OCR",
            file_path.name, jumlah,
        )
        return True

    halaman = _pdf_page_count(file_path)
    if halaman and jumlah / halaman < _MIN_PDF_CHARS_PER_PAGE:
        logger.warning(
            "PDF '{}' hanya {:.0f} karakter/halaman ({} karakter, {} halaman) — "
            "lapisan teksnya terlalu tipis; mencoba ulang dengan OCR",
            file_path.name, jumlah / halaman, jumlah, halaman,
        )
        return True

    return False


def _build_partition_kwargs(file_path: Path) -> dict[str, Any]:
    ext = file_path.suffix.lower().lstrip(".")
    if ext in _PDF_EXTENSIONS:
        # infer_table_structure omitted — Unstructured passes it internally for PDFs
        # and would raise "multiple values" error if we also pass it.
        return {
            "strategy": "fast",
        }
    # For docx/pptx/xlsx/etc. partition() auto-routes to the format-specific
    # partitioner, which already sets infer_table_structure (default True, so
    # tables still get text_as_html). Passing it here again would raise
    # "got multiple values for keyword argument 'infer_table_structure'".
    return {
        "strategy": "auto",
    }


def _assemble_elements(
    elements: list[Element],
    source_file: str,
    content_id: str | None = None,
) -> list[ParsedElement]:
    """Gabungkan potongan mentah Unstructured jadi ParsedElement.

    Teks yang berurutan disatukan supaya satu paragraf tidak pecah jadi
    beberapa elemen — tapi HANYA selama masih di halaman yang sama. Pergantian
    halaman memaksa penutupan buffer.

    Tanpa penutupan itu, dokumen yang isinya teks semua (tidak ada tabel atau
    gambar yang memaksa flush) akan menyatu jadi SATU elemen dengan nomor
    halaman elemen pertamanya. Terukur: PDF 29 halaman keluar sebagai satu
    elemen "halaman 1", lalu seluruh chunk turunannya mewarisi angka itu.
    Akibatnya nomor halaman tidak menunjuk ke mana pun — sitasi ke mahasiswa
    salah, dan rentang halaman yang dicatat peta konsep tidak membedakan satu
    konsep dari konsep lain.

    Halaman ``None`` (mis. .docx dan .txt yang memang tidak berhalaman) tidak
    dianggap sebagai pergantian: menutup buffer di situ akan memecah teks tanpa
    memberi informasi apa pun sebagai gantinya.
    """
    parsed: list[ParsedElement] = []
    text_buffer: list[str] = []
    buffer_page: int | None = None

    def flush_text() -> None:
        nonlocal text_buffer, buffer_page
        if not text_buffer:
            return
        parsed.append(
            _make_text_element(text_buffer, source_file, buffer_page, content_id)
        )
        text_buffer = []
        buffer_page = None

    for el in elements:
        page = _page_number_of(el)

        if isinstance(el, UTable):
            flush_text()
            parsed.append(_make_table_element(el, source_file, page, content_id))
            continue

        if isinstance(el, UImage):
            flush_text()
            img = _make_image_element(el, source_file, page, content_id)
            if img is not None:
                parsed.append(img)
            continue

        text = (getattr(el, "text", "") or "").strip()
        if not text:
            continue

        if page is not None and buffer_page is not None and page != buffer_page:
            flush_text()

        text_buffer.append(text)
        if buffer_page is None:
            buffer_page = page

    flush_text()
    return parsed


def _page_number_of(element: Element) -> int | None:
    metadata = getattr(element, "metadata", None)
    if metadata is None:
        return None
    return getattr(metadata, "page_number", None)


def _make_text_element(
    buffer: list[str], source_file: str, page_number: int | None,
    content_id: str | None,
) -> ParsedElement:
    return ParsedElement(
        element_id=str(uuid.uuid4()),
        element_type=ElementType.TEXT,
        content="\n\n".join(buffer),
        source_file=source_file,
        page_number=page_number,
        content_id=content_id,
    )


def _make_table_element(
    table: UTable, source_file: str, page_number: int | None,
    content_id: str | None,
) -> ParsedElement:
    raw_html: str | None = None
    metadata = getattr(table, "metadata", None)
    if metadata is not None:
        raw_html = getattr(metadata, "text_as_html", None)

    return ParsedElement(
        element_id=str(uuid.uuid4()),
        element_type=ElementType.TABLE,
        content=table.text or "",
        raw_html=raw_html,
        source_file=source_file,
        page_number=page_number,
        content_id=content_id,
    )


def _make_image_element(
    image: UImage, source_file: str, page_number: int | None,
    content_id: str | None,
) -> ParsedElement | None:
    metadata = getattr(image, "metadata", None)
    image_base64: str | None = None
    if metadata is not None:
        image_base64 = getattr(metadata, "image_base64", None)

    if not image_base64:
        logger.warning(
            "Image on page {} of {} has no base64 payload, skipping",
            page_number, source_file,
        )
        return None

    return ParsedElement(
        element_id=str(uuid.uuid4()),
        element_type=ElementType.IMAGE,
        content=image.text or "",
        image_base64=image_base64,
        source_file=source_file,
        page_number=page_number,
        content_id=content_id,
    )


def _count_by_type(elements: list[ParsedElement]) -> dict[ElementType, int]:
    counts = {t: 0 for t in ElementType}
    for el in elements:
        counts[el.element_type] += 1
    return counts
